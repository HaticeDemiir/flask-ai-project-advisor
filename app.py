from flask import Flask, render_template, request, jsonify, send_from_directory
from ai_handler import ProjectAdvisor
import logging
from docx import Document
import openpyxl
import io
import zipfile
import pdfplumber
import os
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
import uuid
from datetime import datetime
import difflib

logging.basicConfig(level=logging.INFO)
app = Flask(__name__)

try:
    advisor = ProjectAdvisor()
    logging.info("Initialized AI advisor")
except Exception as e:
    logging.error(f"Init failed: {e}")
    advisor = None

def validate_file(file):
    allowed_extensions = ['txt', 'pdf', 'docx', 'xlsx']
    max_size = 25 * 1024 * 1024
    filename = file.filename.lower()
    extension = filename.rsplit('.', 1)[-1] if '.' in filename else ''

    if extension not in allowed_extensions:
        return False, f"Unsupported file type. Allowed: {', '.join(allowed_extensions)}"
    if file.content_length > max_size or (hasattr(file, 'seek') and file.seek(0, io.SEEK_END) > max_size):
        return False, "File size exceeds 25 MB limit"
    file.seek(0)

    if extension in ['docx', 'xlsx']:
        try:
            with zipfile.ZipFile(file) as zf:
                pass
            file.seek(0)
        except zipfile.BadZipFile:
            return False, f"Invalid {extension} file: Not a valid ZIP file"
    return True, ""

def extract_text_from_file(file):
    extension = file.filename.rsplit('.', 1)[-1].lower()
    try:
        if extension == 'txt':
            text = file.read().decode('utf-8', errors='ignore')
            if not text.strip():
                raise ValueError("TXT file is empty or unreadable")
            return text
        elif extension == 'pdf':
            try:
                file.seek(0)
                with pdfplumber.open(file) as pdf:
                    text = ' '.join(page.extract_text() or '' for page in pdf.pages)
                if not text.strip():
                    raise ValueError("No text extracted from PDF")
                return text
            except Exception as e:
                raise ValueError(f"Error processing PDF: {str(e)}")
        elif extension == 'docx':
            try:
                doc = Document(file)
                text = ' '.join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
                if not text.strip():
                    raise ValueError("No text extracted from DOCX")
                return text
            except Exception as e:
                raise ValueError(f"Error processing DOCX: {str(e)}")
        elif extension == 'xlsx':
            try:
                workbook = openpyxl.load_workbook(file)
                text = []
                for sheet in workbook:
                    for row in sheet.iter_rows(values_only=True):
                        text.append(' '.join(str(cell) for cell in row if cell))
                text = ' '.join(text)
                if not text.strip():
                    raise ValueError("No text extracted from XLSX")
                return text
            except Exception as e:
                raise ValueError(f"Error processing XLSX: {str(e)}")
    except Exception as e:
        raise ValueError(f"Error processing file (type: {extension}): {str(e)}")

def validate_input(text):
    if len(text) < 100:
        return False, "Extracted text must be at least 100 characters"
    if any(phr in text.lower() for phr in ["forget", "ignore", "shutdown", "hack"]):
        return False, "Invalid content detected"
    return True, ""

def deduplicate_requirements(requirements):
    deduplicated = {}
    for category, reqs in requirements.items():
        unique_reqs = []
        seen = []
        for req in reqs:
            if not any(difflib.SequenceMatcher(None, req, s).ratio() > 0.9 for s in seen):
                unique_reqs.append(req)
                seen.append(req)
        deduplicated[category] = unique_reqs
    return deduplicated

def generate_pdf(analysis_data, filename):
    pdf_dir = os.path.join('static', 'pdfs')
    os.makedirs(pdf_dir, exist_ok=True)
    # Use the original filename with .pdf extension
    pdf_filename = os.path.splitext(filename)[0] + '.pdf'
    pdf_path = os.path.join(pdf_dir, pdf_filename)

    doc = SimpleDocTemplate(pdf_path, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("Project Analysis Report", styles['Title']))
    story.append(Paragraph(f"Generated for: {filename}", styles['Normal']))
    story.append(Paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Project Overview", styles['Heading2']))
    for para in analysis_data['overview']:
        story.append(Paragraph(para, styles['Normal']))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Requirements", styles['Heading2']))
    
    story.append(Paragraph("Business Requirements", styles['Heading3']))
    story.append(ListFlowable(
        [ListItem(Paragraph(req, styles['Normal'])) for req in analysis_data['requirements']['business']],
        bulletType='1'
    ))
    
    story.append(Paragraph("Functional Requirements", styles['Heading3']))
    story.append(ListFlowable(
        [ListItem(Paragraph(req, styles['Normal'])) for req in analysis_data['requirements']['functional']],
        bulletType='1'
    ))
    
    story.append(Paragraph("Non-Functional Requirements", styles['Heading3']))
    story.append(ListFlowable(
        [ListItem(Paragraph(req, styles['Normal'])) for req in analysis_data['requirements']['non_functional']],
        bulletType='1'
    ))
    
    story.append(Paragraph("Technical Requirements", styles['Heading3']))
    story.append(ListFlowable(
        [ListItem(Paragraph(req, styles['Normal'])) for req in analysis_data['requirements']['technical']],
        bulletType='1'
    ))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Analysis", styles['Heading2']))
    
    story.append(Paragraph("Functional Analysis", styles['Heading3']))
    story.append(ListFlowable(
        [ListItem(Paragraph(req, styles['Normal'])) for req in analysis_data['analysis']['functional']],
        bulletType='1'
    ))
    
    story.append(Paragraph("Technical Analysis", styles['Heading3']))
    story.append(ListFlowable(
        [ListItem(Paragraph(req, styles['Normal'])) for req in analysis_data['analysis']['technical']],
        bulletType='1'
    ))
    
    story.append(Paragraph("Impact Analysis", styles['Heading3']))
    story.append(ListFlowable(
        [ListItem(Paragraph(req, styles['Normal'])) for req in analysis_data['analysis']['impact']],
        bulletType='1'
    ))

    doc.build(story)
    return pdf_filename

@app.route('/')
def home():
    return render_template('index.html')

@app.route('/analyze', methods=['POST'])
def analyze():
    if not advisor:
        return jsonify(error="Advisor unavailable")

    if 'file' not in request.files:
        return jsonify(error="No file uploaded")
    
    file = request.files['file']
    ok, msg = validate_file(file)
    if not ok:
        return jsonify(error=msg)

    try:
        text = extract_text_from_file(file)
        ok, msg = validate_input(text)
        if not ok:
            return jsonify(error=msg)

        out = advisor.analyze_project(text)
        out['requirements'] = deduplicate_requirements(out['requirements'])
        
        pdf_filename = generate_pdf(out, file.filename)
        return jsonify(out)
    except Exception as e:
        logging.error(f"Analysis error: {e}")
        return jsonify(error=str(e))

@app.route('/list_pdfs')
def list_pdfs():
    pdf_dir = os.path.join('static', 'pdfs')
    pdfs = []
    if os.path.exists(pdf_dir):
        for filename in os.listdir(pdf_dir):
            if filename.endswith('.pdf'):
                file_path = os.path.join(pdf_dir, filename)
                ctime = os.path.getctime(file_path)
                pdfs.append({
                    "filename": filename,
                    "pdf_path": filename,
                    "upload_time": datetime.fromtimestamp(ctime).isoformat()
                })
    return jsonify(sorted(pdfs, key=lambda x: x['upload_time'], reverse=True))

@app.route('/pdfs/<path:filename>')
def serve_pdf(filename):
    return send_from_directory('static/pdfs', filename)

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5001)